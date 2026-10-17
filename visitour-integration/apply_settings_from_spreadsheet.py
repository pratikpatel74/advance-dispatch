#!/usr/bin/env python3
"""
apply_settings_from_spreadsheet.py
====================================
Reads the completed 'visitour_engineer_settings.xlsx' and pushes all
VISITOUR fields for each engineer via the Advance Dispatch API.

Run:
    python3 apply_settings_from_spreadsheet.py visitour_engineer_settings.xlsx

Requires: openpyxl  (pip install openpyxl)
"""

import sys, json, urllib.request, urllib.parse, urllib.error

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl not found — run: pip install openpyxl")

API = "http://167.233.160.227:3004"

INT_FIELDS = {
    "Capacity TU1":        "vtCapacity",
    "Capacity TU2":        "vtCapacity2",
    "Speed %":             "vtSpeedPct",
    "Work %":              "vtWorkPct",
    "Max Overtime (mins)": "vtOvertime",
    "Min Radius (mi)":     "vtMinRadius",
    "Max Radius (mi)":     "vtMaxRadius",
    "Max Drive (mi)":      "vtMaxDrive",
}

def api_patch(name: str, data: dict):
    body = json.dumps(data).encode()
    req  = urllib.request.Request(
        f"{API}/api/engineer-config/{urllib.parse.quote(name)}",
        data=body, method="PATCH",
        headers={"Content-Type": "application/json"},
    )
    with urllib.request.urlopen(req, timeout=15) as r:
        return json.loads(r.read())


def main(path: str):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active

    # Read header row (row 3)
    headers = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]

    ok = err = skipped = 0
    for row in ws.iter_rows(min_row=4, values_only=True):
        name  = row[0]
        fm_id = row[1]
        if not name:
            continue
        if not fm_id or str(fm_id).strip().upper() in ("", "NO VISITOUR ID", "NONE"):
            print(f"  — {name:<30} (no VISITOUR ID, skipping)")
            skipped += 1
            continue

        payload = {}
        for ci, hdr in enumerate(headers):
            if hdr in INT_FIELDS:
                val = row[ci]
                if val is not None:
                    try:
                        payload[INT_FIELDS[hdr]] = int(val)
                    except (ValueError, TypeError):
                        pass

        if not payload:
            print(f"  — {name:<30} (no values to set)")
            skipped += 1
            continue

        try:
            result = api_patch(name, payload)
            vt = result.get("visitour", "")
            vt_str = " → VISITOUR ✓" if vt == "pushed" else (f" → VISITOUR ✗ ({vt})" if vt else "")
            print(f"  ✓ {name:<30}{vt_str}")
            ok += 1
        except Exception as e:
            print(f"  ✗ {name:<30} {e}")
            err += 1

    print(f"\nDone — ✓ {ok} updated  ✗ {err} failed  — {skipped} skipped")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("Usage: python3 apply_settings_from_spreadsheet.py <path-to-xlsx>")
    main(sys.argv[1])
